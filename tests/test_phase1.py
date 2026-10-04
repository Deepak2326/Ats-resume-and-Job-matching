"""
tests/test_phase1.py
====================
Smoke test for Phase 1: ``parser.py`` + ``ats_scorer.py``.

Generates a valid single-page PDF in-memory (raw PDF syntax, no third-party
PDF writer needed), runs it through the parser and the ATS scorer, and
asserts the Phase 1 contract end-to-end.

Run from the project root:  python tests/test_phase1.py
"""

from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from ats_scorer import (
    extract_heuristic_features,
    extract_skills,
    find_phones,
    score_resume,
)
from parser import parse_resume

SAMPLE_RESUME_LINES = [
    "JANE DOE",
    "jane.doe@email.com | +1 (415) 555-0132 | linkedin.com/in/janedoe",
    "San Francisco, CA",
    "",
    "SUMMARY",
    "Data scientist with 6 years of experience building ML products.",
    "",
    "SKILLS",
    "Python, SQL, Machine Learning, Pandas, Scikit-learn, AWS, Docker, Tableau, Git",
    "",
    "EXPERIENCE",
    "Senior Data Scientist - Acme Corp (2020-Present)",
    "- Led a team of 5 engineers to deploy a recommendation engine serving 2M+ users.",
    "- Increased revenue by 32% through targeted personalization models.",
    "- Reduced cloud costs by $120K annually via pipeline optimization.",
    "- Built an NLP classifier that improved routing accuracy by 18%.",
    "- Automated weekly reporting, saving 10+ analyst hours per week.",
    "Data Analyst - Beta LLC (2018-2020)",
    "- Developed Tableau dashboards adopted by 3 departments.",
    "- Improved forecast accuracy by 12% using A/B testing and regression.",
    "",
    "EDUCATION",
    "M.S. Computer Science - State University (2018)",
    "",
    "PROJECTS",
    "- Created an open-source ETL library with 500+ GitHub stars.",
    "",
    "CERTIFICATIONS",
    "AWS Certified Machine Learning - Specialty",
]


def build_sample_pdf(lines: list[str]) -> bytes:
    """Build a minimal but valid one-page PDF containing ``lines`` of text."""
    ops = ["BT", "/F1 12 Tf", "14 TL", "72 720 Td"]
    for i, line in enumerate(lines):
        safe = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        ops.append(f"({safe}) Tj" if i == 0 else f"T* ({safe}) Tj")
    ops.append("ET")
    content = "\n".join(ops).encode("latin-1", errors="replace")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n"
        + content + b"\nendstream",
    ]

    pdf = bytearray(b"%PDF-1.4\n")
    offsets = []
    for num, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{num} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n".encode()
    pdf += b"0000000000 65535 f \n"
    for off in offsets:
        pdf += f"{off:010d} 00000 n \n".encode()
    pdf += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF"
    ).encode()
    return bytes(pdf)


def run_international_signal_checks() -> None:
    """Regression: Indian phone formats, LinkedIn/GitHub/portfolio URLs, and
    the bullet glyphs Word-authored documents actually emit."""
    intl_text = (
        "RAVI KUMAR\n"
        "+91 98765 43210 | ravi.k@gmail.com\n"
        "in.linkedin.com/in/ravikumar | github.com/ravikumar | https://ravi.dev\n"
        "EXPERIENCE\n"
        "Senior Data Analyst - FinEdge (2019-2024)\n"
        "\uf0b7 Built ETL pipelines processing 5M+ rows daily\n"
        "\u2014 Deployed ML models reducing churn by 18%\n"
        "\u25cf Automated weekly MIS reports saving $40K annually\n"
        "1. Won the national data hackathon\n"
    )
    feats = extract_heuristic_features(intl_text)
    assert feats["has_phone"], "Indian mobile (+91 98765 43210) not detected"
    assert feats["has_email"], "Email not detected"
    assert feats["has_linkedin"], "Locale LinkedIn URL (in.linkedin.com) not detected"
    assert feats["has_github"], "GitHub URL not detected"
    assert feats["has_portfolio"], "Personal-site portfolio URL not detected"
    assert feats["bullet_count"] >= 4, f"Word-style bullets missed: {feats['bullet_count']}"

    # Accepted phone shapes …
    for ok in ("+91-9876543210", "+919876543210", "098765 43210",
               "9876543210", "(415) 555-0132", "+44 20 7946 0958"):
        assert find_phones(ok), f"Phone format missed: {ok}"
    # … and shapes that must NOT count as phones
    for bad in ("2018-2020", "32% growth", "$120K revenue", "500+ users", "2020"):
        assert not find_phones(bad), f"False-positive phone: {bad}"
    print("[intl] OK -> +91 phone, LinkedIn/GitHub/portfolio URLs, Word bullets detected")


def run_docx_structure_checks() -> None:
    """Regression: Word list bullets and clickable hyperlinks survive DOCX parsing."""
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
    except ImportError:
        print("[docx] SKIPPED -> python-docx not installed")
        return

    document = Document()
    document.add_paragraph("RAVI KUMAR")
    document.add_paragraph("+91 98765 43210 | ravi.k@gmail.com")

    # Clickable hyperlink — the URL lives in the .rels part, not the text.
    para = document.add_paragraph("Profiles: ")
    r_id = para.part.relate_to(
        "https://www.linkedin.com/in/ravikumar", RT.HYPERLINK, is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = "LinkedIn"
    run.append(text_el)
    link.append(run)
    para._p.append(link)

    document.add_paragraph("EXPERIENCE")
    document.add_paragraph("Built ETL pipelines in Python", style="List Bullet")
    document.add_paragraph("Deployed ML models on AWS", style="List Bullet")
    document.add_paragraph("Automated MIS reporting with SQL", style="List Bullet")

    import io as _io

    buffer = _io.BytesIO()
    document.save(buffer)
    resume = parse_resume(buffer.getvalue(), filename="ravi_resume.docx")

    assert resume.text.count("\u2022 ") >= 3, f"Word list bullets lost:\n{resume.text}"
    assert "https://www.linkedin.com/in/ravikumar" in resume.links, resume.links
    assert "linkedin.com/in/ravikumar" in resume.text, "Hyperlink URL not merged into text"
    feats = extract_heuristic_features(resume.text)
    assert feats["has_linkedin"], "LinkedIn not detected from DOCX hyperlink"
    assert feats["has_phone"], "Indian mobile not detected in DOCX"
    assert feats["bullet_count"] >= 3, f"DOCX bullets missed: {feats['bullet_count']}"
    print(f"[docx] OK -> {feats['bullet_count']} bullets, hyperlinks={resume.links}")


def main() -> None:
    pdf_bytes = build_sample_pdf(SAMPLE_RESUME_LINES)
    resume = parse_resume(pdf_bytes, filename="sample_resume.pdf")

    # --- Parser assertions -------------------------------------------------
    assert not resume.is_empty, "Parser returned an empty resume"
    assert resume.word_count > 100, f"Too few words: {resume.word_count}"
    assert "SUMMARY" in resume.text, "Section header lost during extraction"
    print(f"[parser] OK -> {resume.word_count} words, {resume.page_count} page(s)")

    # --- Heuristic assertions ----------------------------------------------
    feats = extract_heuristic_features(resume.text)
    assert feats["has_email"] and feats["has_phone"], "Contact info not detected"
    assert feats["has_linkedin"], "LinkedIn URL not detected"
    assert feats["metric_count"] >= 4, f"Metrics under-counted: {feats['metric_count']}"
    assert feats["bullet_count"] >= 7, f"Bullets under-counted: {feats['bullet_count']}"
    assert feats["sections_present"] >= 5, f"Sections missed: {feats['sections']}"
    print(f"[heuristics] OK -> {feats['sections_present']} sections, "
          f"{feats['action_verb_count']} action verbs, {feats['metric_count']} metrics")

    # --- Skills assertions ---------------------------------------------------
    skills = extract_skills(resume.text)
    skill_names = {s["skill"] for s in skills}
    assert "Python" in skill_names and "SQL" in skill_names, f"Skills off: {skill_names}"
    print(f"[skills] OK -> {', '.join(s['skill'] for s in skills[:8])}")

    # --- Scorer assertions ---------------------------------------------------
    report = score_resume(resume.text)
    assert 0 <= report.overall_score <= 100, f"Score out of range: {report.overall_score}"
    assert report.overall_score >= 65, f"Strong resume scored too low: {report.overall_score}"
    assert report.skills, "Report carries no skills"
    assert report.feedback, "Report carries no feedback"
    print(f"[scorer] OK -> score={report.overall_score}/100 (grade {report.grade}), "
          f"ml={report.ml_score} [{report.model_backend}], heuristic={report.heuristic_score}")

    # --- Regression: parsing fixes (Indian phones, profile URLs, bullets) ---
    run_international_signal_checks()
    run_docx_structure_checks()

    print("\nALL PHASE 1 SMOKE TESTS PASSED")


if __name__ == "__main__":
    main()
