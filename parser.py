"""
parser.py
=========
Resume text-extraction utilities for the ATS Resume Analyzer (Phase 1).

Supports:

* **PDF** — via ``pdfplumber`` (page-order preserving text/layout extraction,
  plus hyperlink annotation targets).
* **DOCX** — via ``python-docx`` (preserves Word list bullets as "• " prefixes
  and harvests embedded hyperlink URLs), with a ``docx2txt`` fallback.

The public entry point is :func:`parse_resume`, which accepts a filesystem
path, raw ``bytes``, or a file-like object (e.g. a Streamlit ``UploadedFile``)
and returns a normalised :class:`ParsedResume` dataclass that downstream
engines (ATS scorer, semantic matcher) can consume without caring about the
original file format.
"""

from __future__ import annotations

import io
import logging
import os
import re
from dataclasses import dataclass, field
from typing import Any, BinaryIO, Union

import docx2txt
import pdfplumber

try:  # python-docx: preserves Word list formatting and embedded hyperlinks
    import docx

    _PYTHON_DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover - optional but recommended
    docx = None  # type: ignore[assignment]
    _PYTHON_DOCX_AVAILABLE = False

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS: tuple[str, ...] = (".pdf", ".docx")

_WORD_RE = re.compile(r"\b[\w'+-]+\b")

# Accepted sources: a path string, raw bytes, or an open binary file object.
ResumeSource = Union[str, bytes, bytearray, BinaryIO, Any]


class ResumeParserError(Exception):
    """Raised when a resume cannot be parsed at all (bad type, corrupt file)."""


@dataclass
class ParsedResume:
    """Normalised container for extracted resume content."""

    filename: str
    text: str
    file_type: str
    page_count: int = 0
    char_count: int = 0
    word_count: int = 0
    extraction_warnings: list[str] = field(default_factory=list)
    links: list[str] = field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        """True when no machine-readable text could be extracted."""
        return self.word_count == 0


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------
def _normalise_whitespace(text: str) -> str:
    """Collapse odd whitespace while preserving paragraph boundaries."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_bytes(source: ResumeSource) -> tuple[bytes, str]:
    """Normalise any supported input into ``(raw_bytes, filename)``."""
    if hasattr(source, "read"):  # file-like (Streamlit UploadedFile, open())
        name = getattr(source, "name", "uploaded_resume.pdf")
        if hasattr(source, "seek"):
            source.seek(0)
        data = source.read()
        return (data if isinstance(data, bytes) else data.encode()), name
    if isinstance(source, (bytes, bytearray)):
        return bytes(source), "resume.bin"
    with open(source, "rb") as fh:  # assume path-like
        return fh.read(), os.path.basename(str(source))


# ---------------------------------------------------------------------------
# Format-specific extractors
# ---------------------------------------------------------------------------
def extract_text_from_pdf(source: ResumeSource) -> tuple[str, int, list[str], list[str]]:
    """Extract text from a PDF using pdfplumber.

    Returns ``(text, page_count, warnings, links)``. Pages that contain no
    extractable text (e.g. scanned images) are flagged in ``warnings`` so the
    UI can suggest OCR instead of failing silently. ``links`` collects
    hyperlink annotation targets — clickable URLs live in the PDF's
    annotation layer, not the text layer, so plain extraction never sees them.
    """
    raw, _ = _read_bytes(source)
    warnings: list[str] = []
    links: list[str] = []
    pages_text: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            page_count = len(pdf.pages)
            for i, page in enumerate(pdf.pages, start=1):
                try:
                    page_text = page.extract_text() or ""
                    for annot in (page.annots or []):
                        uri = (
                            annot.get("uri")
                            or (annot.get("data") or {}).get("URI")
                            or ""
                        )
                        if isinstance(uri, str) and uri.strip():
                            links.append(uri.strip())
                except Exception as exc:  # layout edge cases per page
                    warnings.append(f"Page {i}: text extraction failed ({exc}).")
                    page_text = ""
                if not page_text.strip():
                    warnings.append(
                        f"Page {i}: no extractable text (possibly a scanned image)."
                    )
                pages_text.append(page_text)
    except Exception as exc:
        raise ResumeParserError(f"Unable to read PDF: {exc}") from exc
    return _normalise_whitespace("\n\n".join(pages_text)), page_count, warnings, links


def _docx_paragraph_is_list_item(para: Any) -> bool:
    """True when a paragraph carries Word list formatting.

    Word stores auto-bullets/numbering as *formatting* (``numPr`` referencing
    ``numbering.xml``) or via "List *" styles — never as text characters —
    which is why plain text extractors lose them entirely.
    """
    pPr = getattr(para._p, "pPr", None)
    if pPr is not None and getattr(pPr, "numPr", None) is not None:
        return True
    try:
        return (para.style.name or "").lower().startswith("list")
    except Exception:  # broken/missing style definition
        return False


def _docx_hyperlinks(document: Any) -> list[str]:
    """Collect external hyperlink targets from every DOCX part (body, headers,
    footers). URLs live in relationship files (``*.rels``), not in the text."""
    urls: list[str] = []
    try:
        parts = list(document.part.package.iter_parts())
    except Exception:
        parts = [document.part]
    for part in parts:
        for rel in getattr(part, "rels", {}).values():
            if getattr(rel, "is_external", False) and rel.reltype.endswith("/hyperlink"):
                urls.append(rel.target_ref)
    return urls


def _docx_block_lines(document: Any) -> list[str]:
    """Body-order text lines (paragraphs + tables); Word list items are
    re-prefixed with a bullet glyph so bullet-point detection works."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    lines: list[str] = []

    def add_paragraph(para: Paragraph) -> None:
        text = para.text.strip()
        if not text:
            return
        prefix = "\u2022 " if _docx_paragraph_is_list_item(para) else ""
        lines.append(prefix + text)

    for block in document.element.body:  # preserves paragraph/table order
        if block.tag.endswith("}p"):
            add_paragraph(Paragraph(block, document))
        elif block.tag.endswith("}tbl"):
            for row in Table(block, document).rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        add_paragraph(para)
    return lines


def extract_text_from_docx(source: ResumeSource) -> tuple[str, int, list[str], list[str]]:
    """Extract text from a DOCX. Returns ``(text, 1, warnings, links)``.

    Primary path: python-docx — preserves list formatting (as "• " prefixes)
    and harvests embedded hyperlink URLs. Falls back to docx2txt when
    python-docx is unavailable or fails.
    """
    raw, _ = _read_bytes(source)
    warnings: list[str] = []
    links: list[str] = []
    text = ""
    if _PYTHON_DOCX_AVAILABLE:
        try:
            document = docx.Document(io.BytesIO(raw))
            text = "\n".join(_docx_block_lines(document))
            links = _docx_hyperlinks(document)
        except Exception as exc:
            warnings.append(
                f"python-docx extraction failed ({exc}); used docx2txt fallback."
            )
            text = ""
    else:
        warnings.append(
            "python-docx unavailable; bullet formatting and hyperlinks may be lost."
        )
    if not text:
        try:
            text = docx2txt.process(io.BytesIO(raw)) or ""
        except Exception as exc:
            raise ResumeParserError(f"Unable to read DOCX: {exc}") from exc
    return _normalise_whitespace(text), 1, warnings, links


# ---------------------------------------------------------------------------
# Public dispatcher
# ---------------------------------------------------------------------------
def parse_resume(source: ResumeSource, filename: str | None = None) -> ParsedResume:
    """Parse a resume from any supported source into a :class:`ParsedResume`.

    ``filename`` is optional and only needed when ``source`` is raw bytes or an
    unnamed file-like object (the extension drives format detection).
    """
    raw, inferred_name = _read_bytes(source)
    name = filename or inferred_name
    ext = os.path.splitext(name)[1].lower()

    if ext == ".pdf":
        text, pages, warnings, links = extract_text_from_pdf(raw)
        file_type = "pdf"
    elif ext == ".docx":
        text, pages, warnings, links = extract_text_from_docx(raw)
        file_type = "docx"
    else:
        raise ResumeParserError(
            f"Unsupported file type '{ext or 'unknown'}'. "
            f"Supported formats: {', '.join(SUPPORTED_EXTENSIONS)}."
        )

    # Hyperlinks live outside the text layer (DOCX relationships / PDF link
    # annotations). Append the deduplicated targets so downstream ATS signals
    # (LinkedIn / GitHub / portfolio detection) can see them.
    links = sorted(set(links))
    if links:
        text = (text + "\n\n" + "\n".join(links)).strip()

    resume = ParsedResume(
        filename=name,
        text=text,
        file_type=file_type,
        page_count=pages,
        char_count=len(text),
        word_count=len(_WORD_RE.findall(text)),
        extraction_warnings=warnings,
        links=links,
    )
    logger.info(
        "Parsed %s (%s): %d words across %d page(s), %d warning(s).",
        name, file_type, resume.word_count, pages, len(warnings),
    )
    return resume
